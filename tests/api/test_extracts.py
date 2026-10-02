"""Extracts and personal API tokens."""

import csv
import io
import json
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook

from app.core import api_tokens
from app.core.api_tokens import new_token
from app.core.security import decode_token, require_admin
from app.main import app
from app.models.asset import Asset
from app.models.extract import ApiToken
from app.models.remediation import RemediationAction
from app.models.vulnerability import AssetVulnerability, Status, Vulnerability
from app.parsers.utils import remediation
from app.services.ingestion import ingest_findings
from tests.conftest import _make_user


def finding(ip, cve, cvss=8.0, hostname=None):
    return {
        "ip_address": ip,
        "hostname": hostname,
        "operating_system": None,
        "cve_id": cve,
        "title": f"Title of {cve}",
        "description": None,
        "cvss_score": cvss,
        "severity": "High",
        "remediations": [remediation("kb", "KB5034127")],
    }


@pytest.fixture
def backlog(db_session):
    db_session.add(Asset(ip_address="10.0.0.1", owner_team="Servers"))
    db_session.commit()
    ingest_findings(
        db_session,
        [
            finding("10.0.0.1", "CVE-2024-0001", 9.0, hostname="srv-a"),
            finding("10.0.0.2", "CVE-2024-0002", 5.0, hostname="=HYPERLINK(1)"),
        ],
        "nessus",
    )
    vuln = db_session.query(Vulnerability).filter_by(cve_id="CVE-2024-0001").one()
    vuln.in_kev = True
    db_session.commit()


def rows(response):
    assert response.status_code == 200, response.text
    return list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))


@pytest.fixture
def act_as(client, db_session):
    def switch(role):
        user = _make_user(db_session, f"x-{role}", role)
        app.dependency_overrides.pop(require_admin, None)
        app.dependency_overrides[decode_token] = lambda: {
            "sub": str(user.id),
            "role": role,
            "username": user.username,
        }
        return user

    return switch


@pytest.fixture
def token_for(db_session):
    """A real personal token for a user of ``role``, and its secret."""

    def make(role="analyst", lifetime=timedelta(days=30)):
        user = _make_user(db_session, f"pat-{role}", role)
        token, secret = new_token(db_session, user, "script", lifetime)
        db_session.commit()
        return user, token, secret

    return make


def bearer(secret):
    return {"Authorization": f"Bearer {secret}"}


def real_auth():
    """Drop the client fixture's identity: tokens must be checked for real."""
    app.dependency_overrides.pop(decode_token, None)
    app.dependency_overrides.pop(require_admin, None)


class TestExtracts:
    def test_the_backlog_worst_first(self, client, backlog):
        data = rows(client.get("/api/v1/extracts/findings"))

        assert [row["cve_id"] for row in data] == ["CVE-2024-0001", "CVE-2024-0002"]
        assert data[0]["remediation"] == "KB5034127"
        assert data[0]["owner_team"] == "Servers"
        assert data[0]["in_kev"] == "yes"

    def test_columns_filters_and_limit(self, client, backlog):
        response = client.get(
            "/api/v1/extracts/findings",
            params={"columns": "cve_id,risk_score", "kev_only": "true"},
        )
        assert rows(response) == [{"cve_id": "CVE-2024-0001", "risk_score": "9.0"}]

        response = client.get(
            "/api/v1/extracts/findings", params={"columns": "cve_id", "limit": 1}
        )
        assert len(rows(response)) == 1

    def test_json(self, client, backlog):
        response = client.get(
            "/api/v1/extracts/findings",
            params={"format": "json", "columns": "cve_id,in_kev,detected_at"},
        )

        assert response.headers["content-type"].startswith("application/json")
        data = json.loads(response.content)
        assert data[0]["cve_id"] == "CVE-2024-0001"
        assert data[0]["in_kev"] is True
        assert data[0]["detected_at"].startswith("20")

    def test_xlsx(self, client, backlog):
        response = client.get(
            "/api/v1/extracts/findings",
            params={"format": "xlsx", "columns": "cve_id,risk_score,in_kev,asset"},
        )

        assert response.status_code == 200, response.text
        assert response.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert ".xlsx" in response.headers["content-disposition"]
        sheet = load_workbook(io.BytesIO(response.content)).active
        values = [[cell.value for cell in row] for row in sheet.iter_rows()]
        assert values[0] == ["cve_id", "risk_score", "in_kev", "asset"]
        assert values[1][:3] == ["CVE-2024-0001", 9.0, "yes"]
        # A number stays a number, and a text is never evaluated: no quote.
        assert sheet["B2"].data_type == "n"
        assert "=HYPERLINK(1)" in [row[3] for row in values]

    def test_spreadsheet_formulas_are_neutralized(self, client, backlog):
        data = rows(client.get("/api/v1/extracts/findings", params={"columns": "asset"}))
        assert "'=HYPERLINK(1)" in [row["asset"] for row in data]

    @pytest.mark.parametrize(
        "params",
        [
            {"nope": "1"},
            {"kev_only": "maybe"},
            {"min_risk": "11"},
            {"status": "Closed"},
            {"columns": "cve_id,nope"},
            {"format": "pdf"},
        ],
    )
    def test_bad_parameters_are_refused(self, client, backlog, params):
        """A mistyped filter must not silently export everything."""
        response = client.get("/api/v1/extracts/findings", params=params)
        assert response.status_code == 422

    def test_other_datasets(self, client, db_session, backlog):
        assets = {
            row["ip_address"]: row for row in rows(client.get("/api/v1/extracts/assets"))
        }
        assert assets["10.0.0.1"]["open_findings"] == "1"

        [fix] = rows(client.get("/api/v1/extracts/remediation_actions"))
        assert (fix["reference"], fix["hosts"], fix["findings"]) == (
            "KB5034127",
            "2",
            "2",
        )

        cves = rows(
            client.get("/api/v1/extracts/vulnerabilities", params={"kev_only": "yes"})
        )
        assert [row["cve_id"] for row in cves] == ["CVE-2024-0001"]

        action = db_session.query(RemediationAction).one()
        client.post(f"/api/v1/remediation/actions/{action.id}/tickets")
        tickets = rows(client.get("/api/v1/extracts/tickets", params={"status": "all"}))
        assert {row["owner_team"] for row in tickets} == {"Servers", ""}

    def test_an_unknown_dataset_is_404(self, client):
        assert client.get("/api/v1/extracts/nope").status_code == 404

    def test_datasets_follow_the_modules(self, client, act_as):
        act_as("remediator")
        keys = [d["key"] for d in client.get("/api/v1/extracts/datasets").json()]

        assert keys == ["findings", "remediation_actions", "tickets", "assets"]
        assert client.get("/api/v1/extracts/vulnerabilities").status_code == 403


class TestSavedExtracts:
    def test_save_and_run(self, client, backlog):
        response = client.post(
            "/api/v1/extracts/saved",
            json={
                "name": "KEV backlog",
                "dataset": "findings",
                "columns": ["cve_id", "asset"],
                "filters": {"kev_only": True},
                "format": "csv",
            },
        )
        assert response.status_code == 201
        saved = response.json()
        assert saved["run_path"] == f"/api/v1/extracts/saved/{saved['id']}/run"

        assert rows(client.get(saved["run_path"])) == [
            {"cve_id": "CVE-2024-0001", "asset": "srv-a"}
        ]
        as_json = client.get(saved["run_path"], params={"format": "json"})
        assert json.loads(as_json.content)[0]["cve_id"] == "CVE-2024-0001"
        as_xlsx = client.get(saved["run_path"], params={"format": "xlsx"})
        sheet = load_workbook(io.BytesIO(as_xlsx.content)).active
        assert sheet["A2"].value == "CVE-2024-0001"

    def test_saved_as_xlsx(self, client, backlog):
        response = client.post(
            "/api/v1/extracts/saved",
            json={"name": "For Excel", "dataset": "findings", "format": "xlsx"},
        )
        assert response.status_code == 201, response.text

        run = client.get(response.json()["run_path"])

        assert run.headers["content-type"].endswith("spreadsheetml.sheet")
        assert load_workbook(io.BytesIO(run.content)).active.max_row == 3

    def test_invalid_extracts_are_not_saved(self, client):
        response = client.post(
            "/api/v1/extracts/saved",
            json={"name": "x", "dataset": "findings", "filters": {"nope": 1}},
        )
        assert response.status_code == 422

    def test_they_belong_to_their_author(self, client, act_as, backlog):
        saved = client.post(
            "/api/v1/extracts/saved", json={"name": "mine", "dataset": "assets"}
        ).json()
        act_as("analyst")

        assert client.get("/api/v1/extracts/saved").json() == []
        assert client.get(saved["run_path"]).status_code == 404
        assert client.delete(f"/api/v1/extracts/saved/{saved['id']}").status_code == 404


class TestTokens:
    def test_created_once_listed_without_its_secret(self, client):
        response = client.post(
            "/api/v1/extracts/tokens", json={"name": "Power BI", "expires_in_days": 30}
        )

        assert response.status_code == 201
        created = response.json()
        assert created["token"].startswith("vigie_pat_")
        assert created["token"].startswith(created["prefix"])
        assert created["active"] is True
        [listed] = client.get("/api/v1/extracts/tokens").json()
        assert "token" not in listed
        assert listed["name"] == "Power BI"

    @pytest.mark.parametrize(
        "body", [{"name": " "}, {"name": "x", "expires_in_days": 400}]
    )
    def test_invalid_requests(self, client, body):
        assert client.post("/api/v1/extracts/tokens", json=body).status_code == 422

    def test_a_script_reads_with_it(self, unauthenticated_client, token_for, backlog):
        _, token, secret = token_for()

        response = unauthenticated_client.get(
            "/api/v1/extracts/findings", headers=bearer(secret)
        )

        assert len(rows(response)) == 2
        assert (
            unauthenticated_client.get(
                "/api/v1/vulnerabilities/findings", headers=bearer(secret)
            ).status_code
            == 200
        )

    def test_it_records_its_last_use(self, unauthenticated_client, token_for, db_session):
        _, token, secret = token_for()
        unauthenticated_client.get("/api/v1/me/modules", headers=bearer(secret))

        db_session.refresh(token)
        assert token.last_used_at is not None

    def test_it_cannot_change_anything(
        self, unauthenticated_client, token_for, backlog, db_session
    ):
        _, _, secret = token_for()
        finding = db_session.query(AssetVulnerability).first()

        response = unauthenticated_client.patch(
            f"/api/v1/vulnerabilities/findings/{finding.id}",
            json={"status": "Remediated"},
            headers=bearer(secret),
        )

        assert response.status_code == 403
        db_session.refresh(finding)
        assert finding.status == Status.open
        # Nor mint another token.
        response = unauthenticated_client.post(
            "/api/v1/extracts/tokens", json={"name": "x"}, headers=bearer(secret)
        )
        assert response.status_code == 403

    def test_revoked_expired_or_unknown_is_refused(
        self, unauthenticated_client, token_for, db_session
    ):
        _, revoked, revoked_secret = token_for("analyst")
        revoked.revoked_at = datetime.now(UTC)
        _, _, expired_secret = token_for("admin", lifetime=timedelta(seconds=-1))
        db_session.commit()

        for secret in (revoked_secret, expired_secret, "vigie_pat_" + "x" * 43):
            response = unauthenticated_client.get(
                "/api/v1/me/modules", headers=bearer(secret)
            )
            assert response.status_code == 401

    def test_it_stops_with_the_module(self, client, unauthenticated_client, token_for):
        _, _, secret = token_for("analyst")
        client.patch("/api/v1/admin/modules/extracts", json={"enabled": False})
        real_auth()

        response = unauthenticated_client.get(
            "/api/v1/me/modules", headers=bearer(secret)
        )

        assert response.status_code == 403

    def test_its_owner_revokes_it(self, client, unauthenticated_client, db_session):
        created = client.post("/api/v1/extracts/tokens", json={"name": "cron"}).json()

        assert (
            client.delete(f"/api/v1/extracts/tokens/{created['id']}").status_code == 204
        )
        real_auth()
        response = unauthenticated_client.get(
            "/api/v1/me/modules", headers=bearer(created["token"])
        )
        assert response.status_code == 401

    def test_someone_else_cannot(self, client, act_as):
        created = client.post("/api/v1/extracts/tokens", json={"name": "cron"}).json()
        act_as("analyst")

        assert (
            client.delete(f"/api/v1/extracts/tokens/{created['id']}").status_code == 404
        )

    def test_a_ceiling_per_user(self, client, monkeypatch):
        monkeypatch.setattr(api_tokens, "MAX_TOKENS_PER_USER", 1)
        monkeypatch.setattr("app.api.v1.extracts.MAX_TOKENS_PER_USER", 1)
        client.post("/api/v1/extracts/tokens", json={"name": "one"})

        response = client.post("/api/v1/extracts/tokens", json={"name": "two"})

        assert response.status_code == 409

    def test_only_its_hash_is_stored(self, client, db_session):
        created = client.post("/api/v1/extracts/tokens", json={"name": "x"}).json()
        stored = db_session.get(ApiToken, created["id"])
        assert created["token"] not in (stored.token_hash, stored.prefix)
        assert len(stored.token_hash) == 64


class TestAssetTagFilter:
    def test_the_extract_filters_tags_like_the_screen(self, client, db_session):
        db_session.add_all(
            [
                Asset(ip_address="10.50.0.1", tags=["pci_dss"]),
                Asset(ip_address="10.50.0.2", tags=["pciXdss"]),
            ]
        )
        db_session.commit()

        response = client.get(
            "/api/v1/extracts/assets",
            params={"format": "json", "columns": "ip_address", "tag": "PCI_DSS"},
        )

        assert response.status_code == 200, response.text
        assert [row["ip_address"] for row in json.loads(response.content)] == [
            "10.50.0.1"
        ]
