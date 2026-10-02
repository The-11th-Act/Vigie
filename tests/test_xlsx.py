"""The XLSX writer, read back by openpyxl (a test dependency only)."""

import io
import zipfile
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.services.xlsx import MAX_CELL_CHARS, column_letters, sheet_name, write_xlsx


def written(header, rows, name="Findings"):
    buffer = io.BytesIO()
    write_xlsx(buffer, header, rows, name=name)
    buffer.seek(0)
    return load_workbook(buffer)


@pytest.mark.parametrize(
    "index, letters",
    [(0, "A"), (25, "Z"), (26, "AA"), (51, "AZ"), (701, "ZZ"), (702, "AAA")],
)
def test_column_letters(index, letters):
    assert column_letters(index) == letters


def test_a_header_then_typed_cells():
    sheet = written(
        ["cve_id", "risk_score", "count", "price", "in_kev", "note"],
        [["CVE-2024-0001", 9.5, 3, Decimal("1.50"), True, None]],
    ).active

    assert sheet.title == "Findings"
    assert [cell.value for cell in sheet[1]] == [
        "cve_id",
        "risk_score",
        "count",
        "price",
        "in_kev",
        "note",
    ]
    assert all(cell.font.b for cell in sheet[1])
    assert [cell.value for cell in sheet[2]] == [
        "CVE-2024-0001",
        9.5,
        3,
        1.5,
        "yes",
        None,
    ]
    assert sheet["B2"].data_type == "n"
    # The header stays in view.
    assert sheet.freeze_panes == "A2"


def test_text_is_never_a_formula():
    """Unlike CSV, an inline string is never evaluated: no quote is added."""
    sheet = written(["asset"], [["=HYPERLINK(1)"], ["@SUM(A1)"]]).active

    assert sheet["A2"].value == "=HYPERLINK(1)"
    assert sheet["A2"].data_type == "s"
    assert sheet["A3"].value == "@SUM(A1)"


def test_what_xml_or_excel_cannot_hold():
    long_text = "x" * (MAX_CELL_CHARS + 10)
    sheet = written(
        ["text"],
        [
            ['a < b & c > d "quoted"'],
            ["plugin\x00output\x07 with\x1b controls"],
            [long_text],
            [float("nan")],
            [float("inf")],
        ],
    ).active

    assert sheet["A2"].value == 'a < b & c > d "quoted"'
    assert sheet["A3"].value == "pluginoutput with controls"
    assert len(sheet["A4"].value) == MAX_CELL_CHARS
    assert sheet["A5"].value is None and sheet["A6"].value is None


def test_rows_are_streamed_into_one_sheet():
    sheet = written(["n"], ([i] for i in range(5000))).active

    assert sheet.max_row == 5001
    assert sheet["A5001"].value == 4999


def test_the_package_parts():
    buffer = io.BytesIO()
    write_xlsx(buffer, ["a"], [["b"]])
    names = set(zipfile.ZipFile(buffer).namelist())

    assert names == {
        "[Content_Types].xml",
        "_rels/.rels",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/styles.xml",
        "xl/worksheets/sheet1.xml",
    }


@pytest.mark.parametrize(
    "name, expected",
    [
        ("Findings", "Findings"),
        ("Fixes [KB]: a/b?", "Fixes  KB   a b"),
        ("x" * 40, "x" * 31),
        ("", "Sheet1"),
    ],
)
def test_sheet_names(name, expected):
    assert sheet_name(name) == expected
